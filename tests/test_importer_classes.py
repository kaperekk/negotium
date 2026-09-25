"""
test_importer_classes.py — tests for BaseBrokerImporter implementations and ImportService.
"""
from __future__ import annotations

import json
from datetime import date
from pathlib import Path

import openpyxl
import pytest

from domain.models import LedgerEntry, Transaction
from services.import_service import ImportService
from services.importers import (
    BaseBrokerImporter,
    BossaImporter,
    ImportResult,
    ManualImporter,
    ValidationResult,
    XtbImporter,
    ingest_transactions,
)


def test_validation_and_import_result_conversions():
    """ValidationResult and ImportResult properly convert to expected types."""
    v = ValidationResult(valid=True, message="OK")
    assert v.to_tuple() == (True, "OK")

    res = ImportResult(success=True, imported=5, skipped=2, error="warn")
    d = res.to_dict()
    assert d == {"success": True, "imported": 5, "skipped": 2, "error": "warn"}

    res_clean = ImportResult(success=True, imported=3, skipped=0)
    assert res_clean.to_dict() == {"success": True, "imported": 3, "skipped": 0}


def test_import_result_summarises_warnings():
    """Warnings ride along with a successful result and are capped for display."""
    res = ImportResult(success=True, imported=1, warnings=[f"w{i}" for i in range(25)])
    d = res.to_dict()
    assert len(d["warnings"]) == 11  # 10 shown + 1 summary line
    assert "15 more warnings" in d["warnings"][-1]

    assert ImportResult(success=True, imported=1, warnings=["only one"]).to_dict()["warnings"] == ["only one"]
    assert "warnings" not in ImportResult(success=True, imported=1).to_dict()


def test_ingest_transactions_with_transaction_objects(tmp: Path):
    """ingest_transactions handles Transaction models directly."""
    from ledger_core import get_all_transactions

    tx1 = Transaction(
        date="2026-03-01",
        entries=[
            LedgerEntry(ticker="AAPL", amount=10.0),
            LedgerEntry(ticker="USD", amount=-1500.0),
        ],
    )
    tx2 = Transaction(
        date="2026-03-02",
        entries=[
            LedgerEntry(ticker="USD", amount=500.0, account_operation=True),
        ],
    )

    result = ingest_transactions([tx1, tx2])
    assert result.success is True
    assert result.imported == 2
    assert result.skipped == 0

    all_tx = get_all_transactions()
    assert len(all_tx) == 2

    # Ingest again — should deduplicate
    result2 = ingest_transactions([tx1, tx2])
    assert result2.imported == 0
    assert result2.skipped == 2


def test_xtb_importer_class(tmp: Path):
    """XtbImporter class conforms to BaseBrokerImporter interface."""
    importer = XtbImporter()
    assert isinstance(importer, BaseBrokerImporter)

    # Create dummy workbook
    wb = openpyxl.Workbook()
    ws = wb.create_sheet("Cash Operations")
    ws.append(["Type", "Ticker", "Instrument", "Time", "Amount", "ID", "Comment"])
    ws.append(["Deposit", "USD", "", "2026-01-10 10:00:00", 1000.0, 1, ""])
    del wb["Sheet"]
    xlsx_path = tmp / "xtb_test.xlsx"
    wb.save(xlsx_path)
    wb.close()

    val = importer.validate(xlsx_path)
    assert val.valid is True

    parsed = importer.parse(xlsx_path, "USD")
    assert len(parsed.transactions) == 1

    res = importer.import_file(xlsx_path, "USD")
    assert res.success is True
    assert res.imported == 1


def test_xtb_file_currency_from_prefix(tmp: Path):
    """XTB resolves the account currency from the filename prefix, defaulting to EUR."""
    importer = XtbImporter()
    assert importer.file_currency("EUR_history.xlsx") == "EUR"
    assert importer.file_currency("PLN_history.xlsx") == "PLN"
    assert importer.file_currency("2026-09-25_historia.xlsx") == "EUR"


def test_bossa_file_currency_is_never_guessed(tmp: Path):
    """BOSSA reads the currency per row, so no filename prefix may become a ticker."""
    importer = BossaImporter()
    assert importer.file_currency("PLN_bossa.csv") == ""
    assert importer.file_currency("EUR_bossa.csv") == ""
    assert importer.file_currency("historia_finansowa.csv") == ""


def test_bossa_importer_class(tmp: Path):
    """BossaImporter class conforms to BaseBrokerImporter interface."""
    importer = BossaImporter()
    assert isinstance(importer, BaseBrokerImporter)

    csv_path = tmp / "bossa_test.csv"
    csv_path.write_text(
        "data;tytuł operacji;szczegóły;kwota;waluta\n"
        "2026-02-01;Przelew do DM BOŚ;;3000.00;PLN\n",
        encoding="utf-8",
    )

    val = importer.validate(csv_path)
    assert val.valid is True

    parsed = importer.parse(csv_path, "PLN")
    assert len(parsed.transactions) == 1

    res = importer.import_file(csv_path, "PLN")
    assert res.success is True
    assert res.imported == 1


def test_bossa_importer_keeps_unresolved_isins(tmp: Path):
    """The polymorphic path must not drop the unresolved-ISIN report."""
    importer = BossaImporter()
    csv_path = tmp / "bossa_unresolved.csv"
    csv_path.write_text(
        "data;tytuł operacji;szczegóły;kwota;waluta\n"
        "2026-02-01;Rozliczenie transakcji kupna:;"
        "Unknown Fund (DE0005793303) 10 x 100.00 PLN nr 1;-1000,00;PLN\n",
        encoding="utf-8",
    )

    parsed = importer.parse(csv_path)
    assert parsed.transactions == []
    assert "DE0005793303" in parsed.unresolved
    assert any("DE0005793303" in w for w in parsed.warnings)

    res = importer.import_file(csv_path)
    assert res.success is True
    assert res.imported == 0
    assert any("DE0005793303" in w for w in res.warnings)
    assert any("DE0005793303" in w for w in res.to_dict()["warnings"])


def test_manual_importer_class(tmp: Path):
    """ManualImporter class conforms to BaseBrokerImporter interface."""
    importer = ManualImporter()
    assert isinstance(importer, BaseBrokerImporter)

    json_path = tmp / "manual_test.json"
    data = [
        {
            "date": "2026-02-15",
            "entries": [
                {"ticker": "PLN", "amount": 2000.0, "account_operation": True}
            ],
        }
    ]
    json_path.write_text(json.dumps(data), encoding="utf-8")

    val = importer.validate(json_path)
    assert val.valid is True

    parsed = importer.parse(json_path)
    assert len(parsed.transactions) == 1

    res = importer.import_file(json_path)
    assert res.success is True
    assert res.imported == 1


def test_import_service_orchestration(tmp: Path):
    """ImportService scans directories and calls registered polymorphic importers."""
    from storage.context import ProjectContext

    ctx = ProjectContext(name="test_proj", data_root=tmp)
    ctx.ensure_directories()
    xtb_dir = ctx.imports_dir / "xtb"
    bossa_dir = ctx.imports_dir / "bossa"
    custom_dir = ctx.imports_dir / "custom"
    xtb_dir.mkdir(parents=True, exist_ok=True)
    bossa_dir.mkdir(parents=True, exist_ok=True)
    custom_dir.mkdir(parents=True, exist_ok=True)

    # 1. Custom JSON file
    (custom_dir / "manual.json").write_text(
        json.dumps([{"date": "2026-01-01", "entries": [{"ticker": "USD", "amount": 100}]}]),
        encoding="utf-8",
    )

    # 2. BOSSA CSV file
    (bossa_dir / "PLN_bossa.csv").write_text(
        "data;tytuł operacji;szczegóły;kwota;waluta\n"
        "2026-01-02;Przelew do DM BOŚ;;500;PLN\n",
        encoding="utf-8",
    )

    # 3. XTB Excel file
    wb = openpyxl.Workbook()
    ws = wb.create_sheet("Cash Operations")
    ws.append(["Type", "Ticker", "Instrument", "Time", "Amount", "ID", "Comment"])
    ws.append(["Deposit", "EUR", "", "2026-01-03 10:00:00", 250.0, 1, ""])
    del wb["Sheet"]
    wb.save(xtb_dir / "EUR_xtb.xlsx")
    wb.close()

    service = ImportService(context=ctx)
    report = service.run_full_refresh(today=date(2026, 3, 1), base_ccy="PLN")

    assert report.file_count == 3
    assert report.imported == 3
    assert report.warnings == []


def test_import_service_collects_warnings(tmp: Path):
    """Full refresh surfaces per-file warnings instead of swallowing them."""
    from storage.context import ProjectContext

    ctx = ProjectContext(name="warn_proj", data_root=tmp)
    ctx.ensure_directories()
    bossa_dir = ctx.imports_dir / "bossa"
    bossa_dir.mkdir(parents=True, exist_ok=True)
    (bossa_dir / "historia.csv").write_text(
        "data;tytuł operacji;szczegóły;kwota;waluta\n"
        "2026-01-02;Rozliczenie transakcji kupna:;"
        "Unknown Fund (DE0005793303) 10 x 100.00 PLN nr 1;-1000,00;PLN\n",
        encoding="utf-8",
    )

    service = ImportService(context=ctx)
    report = service.run_full_refresh(today=date(2026, 3, 1), base_ccy="PLN")

    assert report.file_count == 1
    assert report.imported == 0
    assert any("historia.csv" in w and "DE0005793303" in w for w in report.warnings)


def test_import_service_single_file_and_dispatch(tmp: Path):
    """import_file routes one file through the registry and resolves its own currency."""
    from storage.context import ProjectContext

    ctx = ProjectContext(name="single_proj", data_root=tmp)
    ctx.ensure_directories()
    bossa_dir = ctx.imports_dir / "bossa"
    bossa_dir.mkdir(parents=True, exist_ok=True)
    csv_path = bossa_dir / "USD_history.csv"
    csv_path.write_text(
        "data;tytuł operacji;szczegóły;kwota;waluta\n"
        "2026-02-01;Przelew do DM BOŚ;;3000.00;USD\n",
        encoding="utf-8",
    )

    service = ImportService(context=ctx)
    assert service.broker_for_filename("anything.csv") == "bossa"
    assert service.broker_for_filename("anything.xlsx") == "xtb"
    assert service.broker_for_filename("anything.txt") is None
    assert service.currency_for("bossa", csv_path) == ""
    assert service.validate_file("bossa", csv_path).valid is True

    result = service.import_file("bossa", csv_path)
    assert result.success is True
    assert result.imported == 1

    with pytest.raises(ValueError):
        service.importer_for("nope")


def test_ingest_warns_on_partially_duplicate_transaction(tmp: Path):
    """A trade whose cash leg is already on file keeps the multiset invariant."""
    from ledger_core import get_all_transactions

    # First import lands both legs.
    tx = Transaction(
        date="2026-04-01",
        entries=[
            LedgerEntry(ticker="IWDA.AS", amount=10.0),
            LedgerEntry(ticker="PLN", amount=-1500.0),
        ],
    )
    assert ingest_transactions([tx]).imported == 1

    # A second statement repeats the cash leg under a different quantity.
    tx2 = Transaction(
        date="2026-04-01",
        entries=[
            LedgerEntry(ticker="IWDA.AS", amount=5.0),
            LedgerEntry(ticker="PLN", amount=-1500.0),
        ],
    )
    result = ingest_transactions([tx2])
    assert result.imported == 1

    entries = get_all_transactions()[0]["entries"]
    assert entries.count({"ticker": "IWDA.AS", "amount": 10.0}) == 1
    assert entries.count({"ticker": "IWDA.AS", "amount": 5.0}) == 1
    assert entries.count({"ticker": "PLN", "amount": -1500.0}) == 1
